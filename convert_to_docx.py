import re
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

def set_cell_background(cell, fill_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(
        f'<w:tcMar {nsdecls("w")}>'
        f'<w:top w:w="{top}" w:type="dxa"/>'
        f'<w:bottom w:w="{bottom}" w:type="dxa"/>'
        f'<w:left w:w="{left}" w:type="dxa"/>'
        f'<w:right w:w="{right}" w:type="dxa"/>'
        f'</w:tcMar>'
    )
    tcPr.append(tcMar)

def create_report_docx():
    doc = docx.Document()

    # Set page margins (1 inch)
    sections = doc.sections
    for section in sections:
        section.top_margin = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin = Inches(1)
        section.right_margin = Inches(1)

    # Base styling
    normal_style = doc.styles['Normal']
    normal_style.font.name = 'Calibri'
    normal_style.font.size = Pt(11)
    normal_style.font.color.rgb = RGBColor(33, 37, 41)
    normal_style.paragraph_format.line_spacing = 1.15
    normal_style.paragraph_format.space_after = Pt(4)

    with open('updated_project_report.md', 'r', encoding='utf-8') as f:
        md_text = f.read()

    lines = md_text.split('\n')
    i = 0
    in_code_block = False
    code_block_lines = []

    while i < len(lines):
        line = lines[i]

        # Handle Code block / Diagram
        if line.strip().startswith('```'):
            if in_code_block:
                in_code_block = False
                # Output code box
                box_text = '\n'.join(code_block_lines)
                table = doc.add_table(rows=1, cols=1)
                table.alignment = WD_TABLE_ALIGNMENT.CENTER
                cell = table.cell(0, 0)
                set_cell_background(cell, "F8F9FA")
                set_cell_margins(cell, 120, 120, 180, 180)
                p = cell.paragraphs[0]
                p.paragraph_format.space_before = Pt(2)
                p.paragraph_format.space_after = Pt(2)
                p.paragraph_format.line_spacing = 1.0
                run = p.add_run(box_text)
                run.font.name = 'Consolas'
                run.font.size = Pt(9.5)
                run.font.color.rgb = RGBColor(40, 40, 40)
                doc.add_paragraph() # Spacing
                code_block_lines = []
            else:
                in_code_block = True
                code_block_lines = []
            i += 1
            continue

        if in_code_block:
            code_block_lines.append(line)
            i += 1
            continue

        # Handle Markdown Table
        if line.strip().startswith('|') and '|' in line.strip()[1:]:
            table_lines = []
            while i < len(lines) and lines[i].strip().startswith('|'):
                table_lines.append(lines[i].strip())
                i += 1
            
            # Parse table
            raw_rows = []
            for tl in table_lines:
                # Skip divider row | :--- | :--- |
                if re.match(r'^\|[\s\:\-]+(?:\|[\s\:\-]+)*\|?$', tl):
                    continue
                cells = [c.strip() for c in tl.strip('|').split('|')]
                raw_rows.append(cells)

            if raw_rows:
                num_cols = max(len(r) for r in raw_rows)
                tbl = doc.add_table(rows=len(raw_rows), cols=num_cols)
                tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
                
                for r_idx, row_cells in enumerate(raw_rows):
                    is_header = (r_idx == 0)
                    for c_idx in range(num_cols):
                        cell = tbl.cell(r_idx, c_idx)
                        text_val = row_cells[c_idx] if c_idx < len(row_cells) else ""
                        set_cell_margins(cell, 80, 80, 120, 120)
                        
                        if is_header:
                            set_cell_background(cell, "1E3A8A") # Navy blue header
                        elif r_idx % 2 == 1:
                            set_cell_background(cell, "F8FAFC") # Light zebra stripe
                        else:
                            set_cell_background(cell, "FFFFFF")

                        p = cell.paragraphs[0]
                        p.paragraph_format.space_before = Pt(2)
                        p.paragraph_format.space_after = Pt(2)
                        p.paragraph_format.line_spacing = 1.0
                        
                        # Clean markdown formatting inside cell
                        clean_text = text_val.replace('**', '').replace('$', '').replace('`', '')
                        run = p.add_run(clean_text)
                        run.font.name = 'Calibri'
                        run.font.size = Pt(9.5)
                        if is_header:
                            run.font.bold = True
                            run.font.color.rgb = RGBColor(255, 255, 255)
                        else:
                            run.font.color.rgb = RGBColor(30, 41, 59)
                
                doc.add_paragraph() # Spacing after table
            continue

        line_str = line.strip()

        # Handle Document Title (# )
        if line_str.startswith('# '):
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(12)
            p.paragraph_format.space_after = Pt(14)
            p.paragraph_format.line_spacing = 1.15
            run = p.add_run(line_str[2:].strip())
            run.font.name = 'Arial'
            run.font.size = Pt(20)
            run.font.bold = True
            run.font.color.rgb = RGBColor(15, 23, 42)
            i += 1
            continue

        # Handle Heading 1 (## )
        if line_str.startswith('## '):
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(14)
            p.paragraph_format.space_after = Pt(4)
            p.paragraph_format.keep_with_next = True
            run = p.add_run(line_str[3:].strip())
            run.font.name = 'Arial'
            run.font.size = Pt(14)
            run.font.bold = True
            run.font.color.rgb = RGBColor(30, 58, 138)
            i += 1
            continue

        # Handle Heading 2 (### )
        if line_str.startswith('### '):
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(10)
            p.paragraph_format.space_after = Pt(3)
            p.paragraph_format.keep_with_next = True
            run = p.add_run(line_str[4:].strip())
            run.font.name = 'Arial'
            run.font.size = Pt(12)
            run.font.bold = True
            run.font.color.rgb = RGBColor(51, 65, 85)
            i += 1
            continue

        # Handle Heading 3 (#### )
        if line_str.startswith('#### '):
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(8)
            p.paragraph_format.space_after = Pt(2)
            p.paragraph_format.keep_with_next = True
            run = p.add_run(line_str[5:].strip())
            run.font.name = 'Calibri'
            run.font.size = Pt(11)
            run.font.bold = True
            run.font.color.rgb = RGBColor(71, 85, 105)
            i += 1
            continue

        # Handle Horizontal Rule (---)
        if line_str in ('---', '***', '___'):
            i += 1
            continue

        # Handle Block Math Equations ($$ ... $$)
        if line_str.startswith('$$') and line_str.endswith('$$') and len(line_str) > 2:
            eq_text = line_str[2:-2].strip()
            # Clean up latex commands for clean Word readability
            clean_eq = (
                eq_text.replace(r'\text{', '')
                       .replace(r'\quad', '   ')
                       .replace(r'\times', ' × ')
                       .replace(r'\le', ' ≤ ')
                       .replace(r'\ge', ' ≥ ')
                       .replace(r'\cdot', ' · ')
                       .replace(r'\sum_{k=1}^K', 'Σ(k=1..K)')
                       .replace(r'\sum_{t=1}^H', 'Σ(t=1..H)')
                       .replace(r'\sum', 'Σ')
                       .replace(r'\sqrt', '√')
                       .replace(r'\mathbf{', '')
                       .replace(r'\boldsymbol{', '')
                       .replace(r'\hat{', '')
                       .replace(r'\Sigma', 'Σ')
                       .replace(r'\mu', 'μ')
                       .replace(r'\sigma', 'σ')
                       .replace(r'\epsilon', 'ε')
                       .replace('}', '')
            )
            p = doc.add_paragraph()
            p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_before = Pt(4)
            p.paragraph_format.space_after = Pt(6)
            run = p.add_run(clean_eq)
            run.font.name = 'Cambria Math'
            run.font.size = Pt(11)
            run.font.italic = True
            run.font.color.rgb = RGBColor(15, 23, 42)
            i += 1
            continue

        # Handle Bullet List (* or -)
        if line_str.startswith('* ') or line_str.startswith('- '):
            content = line_str[2:].strip()
            p = doc.add_paragraph(style='List Bullet')
            p.paragraph_format.space_after = Pt(2)
            
            # Simple bold parsing inside bullets
            parts = re.split(r'(\*\*.*?\*\*)', content)
            for part in parts:
                if part.startswith('**') and part.endswith('**'):
                    r = p.add_run(part[2:-2])
                    r.bold = True
                else:
                    # Clean inline math $...$
                    r_text = part.replace('$', '')
                    p.add_run(r_text)
            i += 1
            continue

        # Handle Numbered List
        num_match = re.match(r'^(\d+)\.\s+(.*)$', line_str)
        if num_match:
            num = num_match.group(1)
            content = num_match.group(2)
            p = doc.add_paragraph(style='List Number')
            p.paragraph_format.space_after = Pt(2)
            parts = re.split(r'(\*\*.*?\*\*)', content)
            for part in parts:
                if part.startswith('**') and part.endswith('**'):
                    r = p.add_run(part[2:-2])
                    r.bold = True
                else:
                    p.add_run(part.replace('$', ''))
            i += 1
            continue

        # Normal Paragraph
        if line_str:
            p = doc.add_paragraph()
            p.paragraph_format.space_after = Pt(4)
            parts = re.split(r'(\*\*.*?\*\*|\*.*?\*)', line_str)
            for part in parts:
                if part.startswith('**') and part.endswith('**'):
                    r = p.add_run(part[2:-2])
                    r.bold = True
                elif part.startswith('*') and part.endswith('*') and len(part) > 2:
                    r = p.add_run(part[1:-1])
                    r.italic = True
                else:
                    p.add_run(part.replace('$', ''))
        
        i += 1

    output_path = 'Multi_Model_Crime_Analytics_Report.docx'
    doc.save(output_path)
    print(f"Successfully generated: {output_path}")

if __name__ == '__main__':
    create_report_docx()
